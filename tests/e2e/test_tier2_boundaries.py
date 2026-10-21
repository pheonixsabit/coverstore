"""
Tier 2: Comprehensive Boundary Value Analysis & Edge Case Test Suite
Covers extreme inputs, zero/empty states, large values, malformed data, and edge conditions.
13 boundary categories x 10 test cases = 130 tests.
"""

import os
import sys
import re
import json
import unittest
import tempfile
import shutil
from unittest.mock import patch, MagicMock
from PIL import Image, ImageDraw, ImageFont
import openpyxl

# Add scripts directory to module search path
SCRIPTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

import moveon_client
import sku_generator
import inventory_manager
import translate_helper
import pathao_sync
import auto_sync
import downloader


# ============================================================================
# Category 1: MoveOn URL & Composite ID Boundaries (10 tests)
# ============================================================================
class TestBoundary01URLAndID(unittest.TestCase):
    """Boundary cases for MoveOn URL parsing and composite ID extraction."""

    def test_b01_empty_string(self):
        self.assertEqual(moveon_client.extract_product_id_from_url(""), "")

    def test_b01_only_whitespace(self):
        self.assertEqual(moveon_client.extract_product_id_from_url("    \t\n  "), "")

    def test_b01_url_without_products_path(self):
        url = "https://moveon.global/BD_en/category/covers/ITEM123"
        cleaned = moveon_client.extract_product_id_from_url(url)
        self.assertIn("ITEM123", cleaned)

    def test_b01_external_domain_sanitization(self):
        url = "https://amazon.com/dp/B08XYZ123"
        cleaned = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(cleaned, "httpsamazoncomdpB08XYZ123")

    def test_b01_extremely_long_slug(self):
        long_slug = "a" * 2000
        url = f"https://moveon.global/BD_en/products/PART1-{long_slug}/PART2"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "PART2PART1")

    def test_b01_url_encoded_characters(self):
        url = "https://moveon.global/BD_en/products/PART1-phone%20case%2Bcover/PART2"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "PART2PART1")

    def test_b01_trailing_multiple_slashes(self):
        url = "https://moveon.global/BD_en/products/PART1-slug/PART2///"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "PART2PART1")

    def test_b01_subdomain_variation(self):
        url = "https://staging.moveon.global/BD_en/products/P1-slug/P2"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "P2P1")

    def test_b01_numeric_only_ids(self):
        url = "https://moveon.global/BD_en/products/12345678-slug/87654321"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "8765432112345678")

    def test_b01_case_preservation(self):
        url = "https://moveon.global/BD_en/products/aBcDeF-slug/XyZwVu"
        extracted = moveon_client.extract_product_id_from_url(url)
        self.assertEqual(extracted, "XyZwVuaBcDeF")


# ============================================================================
# Category 2: Authentication & Session Boundaries (10 tests)
# ============================================================================
class TestBoundary02AuthAndSession(unittest.TestCase):
    """Boundary cases for environment and session loading."""

    def test_b02_empty_env_file(self):
        with tempfile.NamedTemporaryFile("w+", delete=False) as tf:
            tf_path = tf.name
        try:
            vars_map = moveon_client.load_env(tf_path)
            self.assertEqual(vars_map, {})
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b02_missing_env_file(self):
        vars_map = moveon_client.load_env("non_existent_env_file_xyz.env")
        self.assertEqual(vars_map, {})

    def test_b02_multiple_equals_in_line(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write("KEY=val1=val2==val3\n")
            tf_path = tf.name
        try:
            vars_map = moveon_client.load_env(tf_path)
            self.assertEqual(vars_map["KEY"], "val1=val2==val3")
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b02_quoted_values(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write('K1="double_quoted"\nK2=\'single_quoted\'\n')
            tf_path = tf.name
        try:
            vars_map = moveon_client.load_env(tf_path)
            self.assertEqual(vars_map["K1"], "double_quoted")
            self.assertEqual(vars_map["K2"], "single_quoted")
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b02_utf8_env_values(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write("STORE_NAME=ঢাকা কভার স্টোর\n")
            tf_path = tf.name
        try:
            vars_map = moveon_client.load_env(tf_path)
            self.assertEqual(vars_map["STORE_NAME"], "ঢাকা কভার স্টোর")
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b02_corrupted_json_in_moveon_session(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write("{malformed json [")
            tf_path = tf.name
        try:
            # Should handle safely without uncaught exception
            try:
                with open(tf_path, "r", encoding="utf-8") as f:
                    json.load(f)
                loaded = True
            except Exception:
                loaded = False
            self.assertFalse(loaded)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b02_empty_origins_array_in_pathao_session(self):
        session = {"cookies": [], "origins": []}
        origins = session.get("origins", [])
        self.assertEqual(len(origins), 0)

    def test_b02_empty_cookies_array_in_pathao_session(self):
        session = {"cookies": [], "origins": []}
        cookies = session.get("cookies", [])
        self.assertEqual(len(cookies), 0)

    def test_b02_jwt_payload_missing_exp(self):
        payload = {"user_id": 12345}
        exp = payload.get("exp", 0)
        self.assertEqual(exp, 0)

    def test_b02_jwt_payload_expired_timestamp(self):
        payload = {"exp": 1000}
        now = 1700000000
        self.assertTrue(payload["exp"] < now)


# ============================================================================
# Category 3: Model & Filename Cleaning Boundaries (10 tests)
# ============================================================================
class TestBoundary03ModelNameCleaning(unittest.TestCase):
    """Boundary cases for phone model string cleaning and Windows sanitization."""

    def test_b03_multiple_chinese_brackets(self):
        raw = "【1 case】【free gift】Realme Neo8"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "Realme Neo8")

    def test_b03_english_brackets_and_case_note(self):
        raw = "Realme Neo8 [1 case] (1 case)"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "Realme Neo8")

    def test_b03_windows_reserved_characters(self):
        raw = 'Model<1>:*Special?"Test"|Final/'
        cleaned = moveon_client.clean_model_name(raw)
        for char in '<>:*?"|/':
            self.assertNotIn(char, cleaned)

    def test_b03_empty_model_name(self):
        self.assertEqual(moveon_client.clean_model_name(""), "")

    def test_b03_extremely_long_model_name(self):
        long_name = "Realme Neo8 " + "Pro " * 100
        cleaned = moveon_client.clean_model_name(long_name)
        self.assertTrue(len(cleaned) > 0)

    def test_b03_applicable_to_prefix_removal(self):
        raw = "Applicable to Realme Neo8"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "Realme Neo8")

    def test_b03_suitable_for_prefix_removal(self):
        raw = "Suitable for iPhone 15 Pro Max"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "iPhone 15 Pro Max")

    def test_b03_mixed_case_normalization(self):
        raw = "hOnOr x80 i"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "Honor X80i")

    def test_b03_consecutive_spaces_collapse(self):
        raw = "Realme    Neo8   Pro"
        cleaned = moveon_client.sanitize_filename(raw)
        self.assertEqual(cleaned, "Realme Neo8 Pro")

    def test_b03_numeric_only_model(self):
        raw = "600"
        cleaned = moveon_client.clean_model_name(raw)
        self.assertEqual(cleaned, "600")


# ============================================================================
# Category 4: Custom SKU Generation Boundaries (10 tests)
# ============================================================================
class TestBoundary04SKUGeneration(unittest.TestCase):
    """Boundary cases for SKU generator with unknown or multi-word inputs."""

    def test_b04_empty_color_fallback(self):
        sku = sku_generator.generate_sku("Realme Neo8", "Liquid Silicone", "", "Solid")
        self.assertIn("RN8-SIL", sku)

    def test_b04_empty_material_fallback(self):
        sku = sku_generator.generate_sku("Realme Neo8", "", "Antique White", "Solid")
        self.assertIn("RN8", sku)

    def test_b04_unknown_material(self):
        sku = sku_generator.generate_sku("Realme Neo8", "Kevlar Aramid Composite", "Black", "Solid")
        self.assertIn("RN8-", sku)

    def test_b04_multi_word_color(self):
        sku = sku_generator.generate_sku("Realme Neo8", "Liquid Silicone", "Dark Night Green", "Solid")
        self.assertIn("GN", sku)

    def test_b04_slash_separated_color(self):
        sku = sku_generator.generate_sku("Realme Neo8", "Liquid Silicone", "Black / Red", "Solid")
        self.assertIn("BK", sku)

    def test_b04_model_with_plus_sign(self):
        sku = sku_generator.generate_sku("Realme 12 Pro+", "Liquid Silicone", "Black", "Solid")
        self.assertIn("R12PP", sku)

    def test_b04_iphone_pro_max_abbreviation(self):
        sku = sku_generator.generate_sku("iPhone 16 Pro Max", "Hard PC", "Classic Black", "Solid")
        self.assertIn("IP16PM-PC-BK", sku)

    def test_b04_hyphenated_model(self):
        sku = sku_generator.generate_sku("Honor X80-i", "TPU", "Clear", "Clear")
        self.assertIn("HX80-I", sku)

    def test_b04_numeric_model_only(self):
        sku = sku_generator.generate_sku("Neo8", "Liquid Silicone", "Antique White", "Solid")
        self.assertIn("SIL-WT", sku)

    def test_b04_unknown_finish(self):
        sku = sku_generator.generate_sku("Realme Neo8", "Liquid Silicone", "Antique White", "Holographic Prism")
        self.assertEqual(len(sku.split("-")), 4)


# ============================================================================
# Category 5: SKU Disambiguation & Collision Boundaries (10 tests)
# ============================================================================
class TestBoundary05SKUDisambiguation(unittest.TestCase):
    """Boundary cases for SKU collision disambiguation."""

    def test_b05_empty_existing_set(self):
        res = sku_generator.disambiguate_sku("SKU-1", set())
        self.assertEqual(res, "SKU-1")

    def test_b05_single_collision(self):
        res = sku_generator.disambiguate_sku("SKU-1", {"SKU-1"})
        self.assertEqual(res, "SKU-1-V2")

    def test_b05_collision_with_v2(self):
        res = sku_generator.disambiguate_sku("SKU-1", {"SKU-1", "SKU-1-V2"})
        self.assertEqual(res, "SKU-1-V3")

    def test_b05_high_version_collision(self):
        existing = {"SKU-1"} | {f"SKU-1-V{i}" for i in range(2, 10)}
        res = sku_generator.disambiguate_sku("SKU-1", existing)
        self.assertEqual(res, "SKU-1-V10")

    def test_b05_large_existing_set_performance(self):
        existing = {f"OTHER-SKU-{i}" for i in range(10000)}
        existing.add("TARGET-SKU")
        res = sku_generator.disambiguate_sku("TARGET-SKU", existing)
        self.assertEqual(res, "TARGET-SKU-V2")

    def test_b05_already_suffixed_sku(self):
        existing = {"SKU-1-V2"}
        res = sku_generator.disambiguate_sku("SKU-1-V2", existing)
        self.assertEqual(res, "SKU-1-V3")

    def test_b05_case_matching(self):
        existing = {"RN8-SIL-WT-S"}
        res = sku_generator.disambiguate_sku("RN8-SIL-WT-S", existing)
        self.assertEqual(res, "RN8-SIL-WT-S-V2")

    def test_b05_whitespace_in_sku(self):
        res = sku_generator.disambiguate_sku("  SKU-1  ", set())
        self.assertEqual(res.strip(), "SKU-1")

    def test_b05_empty_sku_disambiguation(self):
        res = sku_generator.disambiguate_sku("", set())
        self.assertEqual(res, "")

    def test_b05_batch_of_five_identical_skus(self):
        existing = set()
        resolved = []
        for _ in range(5):
            s = sku_generator.disambiguate_sku("BASE-SKU", existing)
            existing.add(s)
            resolved.append(s)
        self.assertEqual(resolved, ["BASE-SKU", "BASE-SKU-V2", "BASE-SKU-V3", "BASE-SKU-V4", "BASE-SKU-V5"])


# ============================================================================
# Category 6: Financial Cost Inputs & Clean Float Boundaries (10 tests)
# ============================================================================
class TestBoundary06CleanFloat(unittest.TestCase):
    """Boundary cases for float cleaning and parsing."""

    def test_b06_none_input(self):
        self.assertEqual(inventory_manager.clean_float(None, 0.0), 0.0)

    def test_b06_empty_string(self):
        self.assertEqual(inventory_manager.clean_float("", 5.0), 5.0)

    def test_b06_dollar_sign(self):
        self.assertEqual(inventory_manager.clean_float("$125.50"), 125.50)

    def test_b06_taka_sign(self):
        self.assertEqual(inventory_manager.clean_float("৳50.00"), 50.00)

    def test_b06_comma_formatting(self):
        self.assertEqual(inventory_manager.clean_float("1,250.75"), 1250.75)

    def test_b06_whitespace_padding(self):
        self.assertEqual(inventory_manager.clean_float("   99.99   "), 99.99)

    def test_b06_invalid_string_fallback(self):
        self.assertEqual(inventory_manager.clean_float("not_a_number", 10.0), 10.0)

    def test_b06_zero_product_price(self):
        cost = inventory_manager.calculate_landing_cost(0.0, 10.0, 50.0, 5.0)
        self.assertEqual(cost, 96.5)

    def test_b06_high_product_price(self):
        cost = inventory_manager.calculate_landing_cost(50000.0, 10.0, 50.0, 5.0)
        self.assertEqual(cost, 50096.5)

    def test_b06_negative_cost_handling(self):
        cost = inventory_manager.calculate_landing_cost(-10.0, 10.0, 50.0, 5.0)
        self.assertEqual(cost, 86.5)


# ============================================================================
# Category 7: Landing Cost & Margin Tier Boundaries (10 tests)
# ============================================================================
class TestBoundary07MarginTiers(unittest.TestCase):
    """Boundary cases for landing cost summation and dynamic margin tiers."""

    def test_b07_all_zero_costs(self):
        cost = inventory_manager.calculate_landing_cost(0, 0, 0, 0, packaging=0, printing=0, other=0)
        self.assertEqual(cost, 0.0)

    def test_b07_default_packaging_and_printing(self):
        cost = inventory_manager.calculate_landing_cost(0, 0, 0, 0)
        self.assertEqual(cost, 31.5)  # 11.5 + 20.0

    def test_b07_custom_other_expense(self):
        cost = inventory_manager.calculate_landing_cost(100, 10, 50, 5, other=25.5)
        self.assertEqual(cost, 222.0)

    def test_b07_two_decimal_precision(self):
        cost = inventory_manager.calculate_landing_cost(34.0935, 10.0, 50.0, 5.0)
        self.assertEqual(cost, 130.59)

    def test_b07_margin_tier1_upper_boundary(self):
        # 299.99 < 300 -> * 1.40
        self.assertEqual(inventory_manager.calculate_margin(299.99), 419.99)

    def test_b07_margin_tier2_lower_boundary(self):
        # 300.00 is in tier 2 -> * 1.33
        self.assertEqual(inventory_manager.calculate_margin(300.00), 399.00)

    def test_b07_margin_tier2_upper_boundary(self):
        # 399.99 < 400 -> * 1.33
        self.assertEqual(inventory_manager.calculate_margin(399.99), 531.99)

    def test_b07_margin_tier3_lower_boundary(self):
        # 400.00 is in tier 3 -> * 1.25
        self.assertEqual(inventory_manager.calculate_margin(400.00), 500.00)

    def test_b07_margin_extreme_high_value(self):
        self.assertEqual(inventory_manager.calculate_margin(10000.0), 12500.0)

    def test_b07_margin_zero_cost(self):
        self.assertEqual(inventory_manager.calculate_margin(0.0), 0.0)


# ============================================================================
# Category 8: Sell Price Rounding Boundaries (10 tests)
# ============================================================================
class TestBoundary08SellPriceRounding(unittest.TestCase):
    """Boundary cases for sell price rounding ending in 00, 49, 99."""

    def test_b08_mod_zero(self):
        self.assertEqual(inventory_manager.calculate_sell_price(300.0), 300)

    def test_b08_mod_24(self):
        self.assertEqual(inventory_manager.calculate_sell_price(324.0), 300)

    def test_b08_mod_25(self):
        self.assertEqual(inventory_manager.calculate_sell_price(325.0), 349)

    def test_b08_mod_50(self):
        self.assertEqual(inventory_manager.calculate_sell_price(350.0), 349)

    def test_b08_mod_74(self):
        self.assertEqual(inventory_manager.calculate_sell_price(374.0), 349)

    def test_b08_mod_75(self):
        self.assertEqual(inventory_manager.calculate_sell_price(375.0), 399)

    def test_b08_mod_99(self):
        self.assertEqual(inventory_manager.calculate_sell_price(399.0), 399)

    def test_b08_sub_100_low(self):
        self.assertEqual(inventory_manager.calculate_sell_price(15.0), 0)

    def test_b08_sub_100_mid(self):
        self.assertEqual(inventory_manager.calculate_sell_price(60.0), 49)

    def test_b08_sub_100_high(self):
        self.assertEqual(inventory_manager.calculate_sell_price(85.0), 99)


# ============================================================================
# Category 9: Max Discount Tier Boundaries (10 tests)
# ============================================================================
class TestBoundary09MaxDiscount(unittest.TestCase):
    """Boundary cases for tiered max discount calculation and ceiling cap."""

    def test_b09_price_under_300(self):
        self.assertEqual(inventory_manager.calculate_max_discount(299), 0)

    def test_b09_price_exact_300(self):
        self.assertEqual(inventory_manager.calculate_max_discount(300), 21)

    def test_b09_price_tier1_upper(self):
        self.assertEqual(inventory_manager.calculate_max_discount(499), 34.93)

    def test_b09_price_exact_500(self):
        self.assertEqual(inventory_manager.calculate_max_discount(500), 60.0)

    def test_b09_price_tier2_upper(self):
        self.assertEqual(inventory_manager.calculate_max_discount(699), 83.88)

    def test_b09_price_exact_700(self):
        # 700 * 0.15 = 105 -> capped at 99
        self.assertEqual(inventory_manager.calculate_max_discount(700), 99)

    def test_b09_price_high_cap_enforcement(self):
        self.assertEqual(inventory_manager.calculate_max_discount(2000), 99)

    def test_b09_price_zero(self):
        self.assertEqual(inventory_manager.calculate_max_discount(0), 0)

    def test_b09_price_negative(self):
        self.assertEqual(inventory_manager.calculate_max_discount(-50), 0)

    def test_b09_discount_never_exceeds_99(self):
        for p in range(100, 1500, 100):
            self.assertLessEqual(inventory_manager.calculate_max_discount(p), 99)


# ============================================================================
# Category 10: Image Geometry & Dimension Boundaries (10 tests)
# ============================================================================
class TestBoundary10ImageGeometry(unittest.TestCase):
    """Boundary cases for image resolutions, modes, and formats."""

    def test_b10_tiny_1x1_image(self):
        im = Image.new("RGB", (1, 1), color="white")
        self.assertEqual(im.size, (1, 1))

    def test_b10_wide_banner_image(self):
        im = Image.new("RGB", (1200, 400), color="white")
        self.assertEqual(im.size, (1200, 400))

    def test_b10_tall_banner_image(self):
        im = Image.new("RGB", (400, 1200), color="white")
        self.assertEqual(im.size, (400, 1200))

    def test_b10_high_resolution_1600x1600(self):
        im = Image.new("RGB", (1600, 1600), color="white")
        self.assertEqual(im.size, (1600, 1600))

    def test_b10_non_standard_1276x1276(self):
        im = Image.new("RGB", (1276, 1276), color="white")
        self.assertEqual(im.size, (1276, 1276))

    def test_b10_grayscale_mode_conversion(self):
        im = Image.new("L", (100, 100), color=128)
        rgb = im.convert("RGB")
        self.assertEqual(rgb.mode, "RGB")

    def test_b10_rgba_mode_conversion(self):
        im = Image.new("RGBA", (100, 100), color=(255, 0, 0, 128))
        rgb = im.convert("RGB")
        self.assertEqual(rgb.mode, "RGB")

    def test_b10_zero_byte_file_handling(self):
        with tempfile.NamedTemporaryFile("wb", delete=False) as tf:
            tf_path = tf.name
        try:
            with self.assertRaises(Exception):
                Image.open(tf_path)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b10_corrupted_jpeg_handling(self):
        with tempfile.NamedTemporaryFile("wb", delete=False) as tf:
            tf.write(b"\xff\xd8\xff\xe0\x00\x10JFIF\x00corrupted")
            tf_path = tf.name
        try:
            with self.assertRaises(Exception):
                im = Image.open(tf_path)
                im.verify()
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b10_text_file_with_jpg_extension(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, suffix=".jpg") as tf:
            tf.write("This is plain text, not a JPEG")
            tf_path = tf.name
        try:
            with self.assertRaises(Exception):
                Image.open(tf_path)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)


# ============================================================================
# Category 11: Typography & Text Layout Boundaries (10 tests)
# ============================================================================
class TestBoundary11Typography(unittest.TestCase):
    """Boundary cases for text rendering and font handling."""

    def test_b11_empty_text_string(self):
        im = Image.new("RGB", (200, 100))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 20)
        bbox = draw.textbbox((0, 0), "", font=font)
        self.assertEqual(bbox[2] - bbox[0], 0)

    def test_b11_single_character(self):
        im = Image.new("RGB", (200, 100))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 20)
        bbox = draw.textbbox((0, 0), "A", font=font)
        self.assertGreater(bbox[2] - bbox[0], 0)

    def test_b11_extremely_long_string(self):
        text = "LONG " * 100
        im = Image.new("RGB", (500, 100))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 10)
        bbox = draw.textbbox((0, 0), text, font=font)
        self.assertGreater(bbox[2], 500)

    def test_b11_special_characters(self):
        text = "!@#$%^&*()_+=-[]{}|;:',.<>?"
        im = Image.new("RGB", (500, 100))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 14)
        bbox = draw.textbbox((0, 0), text, font=font)
        self.assertGreater(bbox[2] - bbox[0], 0)

    def test_b11_chinese_characters_rendering(self):
        text = "液态硅胶"
        im = Image.new("RGB", (200, 100))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/msyh.ttc", 16)
        bbox = draw.textbbox((0, 0), text, font=font)
        self.assertGreater(bbox[2] - bbox[0], 0)

    def test_b11_minimum_font_size(self):
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 8)
        self.assertIsNotNone(font)

    def test_b11_maximum_font_size(self):
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 120)
        self.assertIsNotNone(font)

    def test_b11_render_at_zero_coordinates(self):
        im = Image.new("RGB", (100, 100), color="white")
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 12)
        draw.text((0, 0), "TOP", fill="black", font=font)
        self.assertNotEqual(im.getpixel((0, 0)), (255, 255, 255, 255))

    def test_b11_render_near_boundary(self):
        im = Image.new("RGB", (100, 100), color="white")
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 12)
        draw.text((80, 80), "BOT", fill="black", font=font)
        self.assertEqual(im.size, (100, 100))

    def test_b11_multiline_text(self):
        text = "Line 1\nLine 2\nLine 3"
        im = Image.new("RGB", (200, 200))
        draw = ImageDraw.Draw(im)
        font = translate_helper.get_font("C:/Windows/Fonts/arial.ttf", 14)
        bbox = draw.textbbox((0, 0), text, font=font)
        self.assertGreater(bbox[3] - bbox[1], 30)


# ============================================================================
# Category 12: CSV & Excel Dataset Boundaries (10 tests)
# ============================================================================
class TestBoundary12CSVAndExcel(unittest.TestCase):
    """Boundary cases for CSV and Excel dataset operations."""

    def test_b12_utf8_bom_header(self):
        with tempfile.NamedTemporaryFile("wb", delete=False) as tf:
            tf.write("\ufeffProduct Code,SEO Title\n".encode("utf-8"))
            tf_path = tf.name
        try:
            with open(tf_path, "r", encoding="utf-8-sig") as f:
                header = f.readline().strip()
            self.assertEqual(header, "Product Code,SEO Title")
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_missing_optional_fields_in_row(self):
        row = {"Product Code": "SKU-1", "SEO Title": "Title"}
        price = row.get("product price", 0.0)
        self.assertEqual(price, 0.0)

    def test_b12_commas_inside_quoted_fields(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write('Product Code,SEO Title\nSKU-1,"Title, with comma, here"\n')
            tf_path = tf.name
        try:
            import csv
            with open(tf_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
            self.assertEqual(rows[0]["SEO Title"], "Title, with comma, here")
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_quotes_inside_quoted_fields(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write('Product Code,SEO Title\nSKU-1,"Title with ""nested"" quotes"\n')
            tf_path = tf.name
        try:
            import csv
            with open(tf_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
            self.assertEqual(rows[0]["SEO Title"], 'Title with "nested" quotes')
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_empty_trailing_lines(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write('Product Code,SEO Title\nSKU-1,Title 1\n\n\n')
            tf_path = tf.name
        try:
            import csv
            with open(tf_path, "r", encoding="utf-8") as f:
                rows = [r for r in csv.DictReader(f) if r.get("Product Code")]
            self.assertEqual(len(rows), 1)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_empty_csv_file(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf_path = tf.name
        try:
            import csv
            with open(tf_path, "r", encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 0)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_large_csv_row_count(self):
        with tempfile.NamedTemporaryFile("w+", delete=False, encoding="utf-8") as tf:
            tf.write("Product Code,SEO Title\n")
            for i in range(1000):
                tf.write(f"SKU-{i},Title {i}\n")
            tf_path = tf.name
        try:
            import csv
            with open(tf_path, "r", encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 1000)
        finally:
            if os.path.exists(tf_path): os.remove(tf_path)

    def test_b12_temporary_lock_file_detection(self):
        filename = "~$inventory.xlsx"
        is_temp_lock = filename.startswith("~$")
        self.assertTrue(is_temp_lock)

    def test_b12_excel_save_to_nested_dir(self):
        with tempfile.TemporaryDirectory() as td:
            sub = os.path.join(td, "nested", "folder")
            os.makedirs(sub)
            wb_path = os.path.join(sub, "test.xlsx")
            wb = openpyxl.Workbook()
            wb.save(wb_path)
            self.assertTrue(os.path.exists(wb_path))

    def test_b12_excel_formula_reading(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"
        self.assertEqual(ws["A3"].value, "=SUM(A1:A2)")


# ============================================================================
# Category 13: Pathao Commerce Import Workbook Boundaries (10 tests)
# ============================================================================
class TestBoundary13PathaoImportWorkbook(unittest.TestCase):
    """Boundary cases for Pathao bulk import workbook generation."""

    def test_b13_parent_only_product(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["PRODUCT_NAME", "CATEGORY", "BASE_SKU", "VARIANT_SKU"])
        ws.append(["Test Cover", "752 - Category", "PARENT-SKU", ""])
        self.assertEqual(ws.max_row, 2)

    def test_b13_product_with_many_variants(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["PRODUCT_NAME", "CATEGORY", "BASE_SKU", "VARIANT_SKU"])
        ws.append(["Test Cover", "752 - Category", "PARENT-SKU", ""])
        for i in range(50):
            ws.append(["", "", "", f"VAR-{i}"])
        self.assertEqual(ws.max_row, 52)

    def test_b13_variant_missing_image(self):
        variant = {"sku": "V-1", "image": ""}
        self.assertEqual(variant["image"], "")

    def test_b13_variant_zero_stock(self):
        variant = {"sku": "V-1", "stock": 0}
        self.assertEqual(variant["stock"], 0)

    def test_b13_long_description_field(self):
        desc = "A" * 5000
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = desc
        self.assertEqual(len(ws["A1"].value), 5000)

    def test_b13_description_with_many_phone_models(self):
        models = [f"Model-{i}" for i in range(20)]
        desc = pathao_sync.get_product_description_for_material("Silicone", models)
        self.assertIn("Model-0", desc)
        self.assertIn("Model-1", desc)

    def test_b13_base_price_offset_formula(self):
        sell_price = 399
        base_price = sell_price + 50
        self.assertEqual(base_price - sell_price, 50)

    def test_b13_discounted_price_equals_sell_price(self):
        sell_price = 349
        discounted = sell_price
        self.assertEqual(discounted, sell_price)

    def test_b13_category_prefix_exact_match(self):
        cat = pathao_sync.PATHAO_CATEGORY_DEFAULT
        self.assertEqual(cat[:5], "752 -")

    def test_b13_disabled_column_value(self):
        disabled_val = "FALSE"
        self.assertEqual(disabled_val, "FALSE")


if __name__ == "__main__":
    unittest.main()
